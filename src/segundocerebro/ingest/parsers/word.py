"""DOCX — 463 files in the corpus, and the format with the cleanest structure.

Word carries real heading styles, so the heading trail is read rather than
guessed. Two details the naive implementation gets wrong:

- `doc.paragraphs` and `doc.tables` are separate collections, so iterating them
  in turn loses document order and every table ends up attached to the last
  heading of the file. Walking the body XML keeps the order.
- Style names are localized. A pt-BR Word writes "Título 1", not "Heading 1",
  and this corpus is full of both.
"""

from __future__ import annotations

import io
import re

from ..document import Block, BlockKind, ParsedDoc
from . import register

NIVEL_DE_ESTILO = re.compile(r"^(?:heading|t[íi]tulo|headline)\s*(\d+)", re.IGNORECASE)
ESTILO_DE_TITULO = re.compile(r"^(?:title|t[íi]tulo)$", re.IGNORECASE)


def _nivel(estilo: str | None) -> int | None:
    """Heading level from a style name, in English or Portuguese."""
    if not estilo:
        return None
    if ESTILO_DE_TITULO.match(estilo.strip()):
        return 1
    m = NIVEL_DE_ESTILO.match(estilo.strip())
    return int(m.group(1)) if m else None


def _itens_na_ordem(documento):  # noqa: ANN001 — tipos internos do python-docx
    """Yield paragraphs and tables in document order."""
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    corpo = documento.element.body
    for filho in corpo.iterchildren():
        if filho.tag == qn("w:p"):
            yield Paragraph(filho, documento)
        elif filho.tag == qn("w:tbl"):
            yield Table(filho, documento)


def _texto_da_tabela(tabela) -> str:  # noqa: ANN001
    linhas: list[str] = []
    for linha in tabela.rows:
        celulas = [c.text.strip() for c in linha.cells]
        # células mescladas repetem o mesmo texto; manter uma vez só
        limpas: list[str] = []
        for c in celulas:
            if c and (not limpas or limpas[-1] != c):
                limpas.append(c)
        if limpas:
            linhas.append(" ; ".join(limpas))
    return "\n".join(linhas)


@register(".docx", ".docm")
def parse_docx(dados: bytes, nome: str) -> ParsedDoc:
    import docx

    documento = docx.Document(io.BytesIO(dados))

    blocos: list[Block] = []
    pilha: list[tuple[int, str]] = []
    corpo: list[str] = []
    n_tabela = 0

    def trilha() -> tuple[str, ...]:
        return tuple(t for _, t in pilha)

    def fechar() -> None:
        conteudo = "\n".join(corpo).strip()
        if conteudo:
            blocos.append(Block(heading_path=trilha(), text=conteudo))
        corpo.clear()

    for item in _itens_na_ordem(documento):
        if hasattr(item, "rows"):  # tabela
            texto = _texto_da_tabela(item)
            if texto:
                fechar()
                n_tabela += 1
                blocos.append(
                    Block(
                        heading_path=trilha(),
                        text=texto,
                        locator=f"tabela {n_tabela}",
                        kind=BlockKind.TABLE,
                    )
                )
            continue

        texto = item.text.strip()
        if not texto:
            continue

        estilo = getattr(getattr(item, "style", None), "name", None)
        nivel = _nivel(estilo)
        if nivel is not None:
            fechar()
            while pilha and pilha[-1][0] >= nivel:
                pilha.pop()
            pilha.append((nivel, texto))
        else:
            corpo.append(texto)

    fechar()

    if not blocos and pilha:
        blocos.append(Block(heading_path=trilha()[:-1], text=pilha[-1][1]))

    meta = {"formato": "docx"}
    propriedades = getattr(documento, "core_properties", None)
    if propriedades is not None:
        if propriedades.title:
            meta["titulo"] = propriedades.title
        if propriedades.author:
            meta["autor"] = propriedades.author

    return ParsedDoc(name=nome, blocks=tuple(blocos), meta=meta)
