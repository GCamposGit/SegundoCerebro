"""PPTX — 251 files, 4,2 GB. The slide is the natural unit.

A slide is already a chunk: it is authored as one idea with a title on top. So
the title becomes the heading and the slide becomes the block, which is both
simpler and better than splitting by size. Speaker notes are kept separate —
they often hold the argument that the slide only gestures at.
"""

from __future__ import annotations

import io

from ..document import Block, BlockKind, ParsedDoc
from . import register


def _texto_do_frame(portador, atributo: str = "text_frame") -> str:  # noqa: ANN001
    """Text of a text frame that may not exist.

    `has_text_frame` não garante que o frame esteja lá: o python-pptx devolve
    `None` quando o `<txBody>` está ausente, e aí `.text` estoura. Medido em
    15/08/2026 no acervo real — 4 decks inteiros ficaram fora do índice com
    `AttributeError: 'NoneType' object has no attribute 'text'`, cada um por
    causa de uma única forma vazia.
    """
    frame = getattr(portador, atributo, None)
    if frame is None:
        return ""
    texto = getattr(frame, "text", None)
    return texto.strip() if isinstance(texto, str) else ""


def _texto_de_forma(forma) -> list[str]:  # noqa: ANN001 — tipos internos do python-pptx
    """Text of a shape, descending into groups and tables."""
    partes: list[str] = []

    if getattr(forma, "has_table", False):
        for linha in forma.table.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" ; ".join(celulas))
        return partes

    if getattr(forma, "shape_type", None) is not None and hasattr(forma, "shapes"):
        for interna in forma.shapes:  # grupo
            partes.extend(_texto_de_forma(interna))
        return partes

    if getattr(forma, "has_text_frame", False):
        texto = _texto_do_frame(forma)
        if texto:
            partes.append(texto)
    return partes


@register(".pptx", ".pptm", version="3")
def parse_pptx(dados: bytes, nome: str) -> ParsedDoc:
    from pptx import Presentation

    apresentacao = Presentation(io.BytesIO(dados))

    titulo_deck = ""
    propriedades = getattr(apresentacao, "core_properties", None)
    if propriedades is not None and propriedades.title:
        titulo_deck = propriedades.title.strip()

    blocos: list[Block] = []
    for numero, slide in enumerate(apresentacao.slides, start=1):
        titulo = ""
        forma_titulo = None
        try:
            forma_titulo = slide.shapes.title
        except (AttributeError, KeyError):
            forma_titulo = None
        if forma_titulo is not None:
            titulo = _texto_do_frame(forma_titulo)

        corpo: list[str] = []
        for forma in slide.shapes:
            if forma is forma_titulo:
                continue
            corpo.extend(_texto_de_forma(forma))

        trilha = tuple(t for t in (titulo_deck, titulo or f"Slide {numero}") if t)
        local = f"slide {numero}"

        conteudo = "\n".join(p for p in corpo if p).strip()
        if conteudo:
            blocos.append(Block(heading_path=trilha, text=conteudo, locator=local))
        elif titulo:
            # slide só de título ainda é recuperável — é o que anuncia a seção
            blocos.append(Block(heading_path=trilha[:-1], text=titulo, locator=local))

        if getattr(slide, "has_notes_slide", False):
            notas = _texto_do_frame(slide.notes_slide, "notes_text_frame")
            if notas:
                blocos.append(
                    Block(heading_path=trilha, text=notas, locator=f"{local} (notas)", kind=BlockKind.SLIDE_NOTES)
                )

    from ..raster_ocr import acrescentar_rasters
    from .ooxml_texto import blocos_de_embeddings, blocos_de_grafico, completar

    completar(dados, blocos)
    blocos.extend(blocos_de_grafico(dados))
    blocos.extend(blocos_de_embeddings(dados, "\n".join(b.text for b in blocos)))
    meta = {"formato": "pptx", "slides": str(len(apresentacao.slides))}
    meta.update(acrescentar_rasters(dados, blocos))
    return ParsedDoc(name=nome, blocks=tuple(blocos), meta=meta)


@register(".ppt", version="4")
def parse_ppt(dados: bytes, nome: str) -> ParsedDoc:
    """PowerPoint 97-2003. Bytes only — no COM, no temp file.

    PPTX com extensão `.ppt` entra aqui, não no despachante: `ooxml` não tem
    parser sem ambiguidade (`parser_for_familia` recusa), e o arquivo virava
    `sem_parser` com o conteúdo à mostra. O tratamento mora no parser da
    extensão que mentiu.
    """
    if dados.startswith(b"PK\x03\x04"):
        doc = parse_pptx(dados, nome)
        meta = dict(doc.meta)
        meta["formato"] = "ppt"
        meta["conteudo_real"] = "pptx"
        return ParsedDoc(name=nome, blocks=doc.blocks, meta=meta)

    from .ole_texto import texto_de_ppt

    texto = texto_de_ppt(dados).strip()
    blocos: tuple[Block, ...] = ()
    if texto:
        blocos = (Block(heading_path=(), text=texto, locator="deck"),)
    return ParsedDoc(name=nome, blocks=blocos, meta={"formato": "ppt"})
