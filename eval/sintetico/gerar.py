"""Gera o corpus sintético de colaboração — empresa fictícia, zero dado real.

    py -m eval.sintetico.gerar

Idempotente: apaga o diretório `corpus/` e reescreve. Os mtimes são parte do
contrato — famílias de versão desempatam por número declarado e, na falta dele,
por data (`retrieve/familias.py`).
"""

from __future__ import annotations

import io
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent
CORPUS = RAIZ / "corpus"

# Datas fixas. A política vigente não declara _vN e é mais nova que a _v6
# (espelho do g010). O deck _v2 é mais velho no mtime que o _v1 (espelho do
# caso em que o número declarado vence a data).
MTIMES = {
    "Politica de IA/PO-VCE-007_Politica de Inteligencia Artificial_v6.docx": "2025-01-15",
    "Politica de IA/PO-VCE-007_Politica de Inteligencia Artificial_revisada_GC.docx": "2026-08-11",
    "Apresentacoes/Deck_governanca_IA_v1.pptx": "2026-06-01",
    "Apresentacoes/Deck_governanca_IA_v2.pptx": "2025-09-01",
    "Propostas/Proposta_Aurora_Tecnica_implantacao_IA.pdf": "2026-02-10",
    "Propostas/Proposta_Boreal_Servicos_implantacao_IA.pdf": "2026-02-12",
    "Contratos/CT-VCE-2024-0142_Servicos_consultoria.pdf": "2024-11-03",
    "Atas/2026-03-12_Ata_mudanca_escopo.md": "2026-03-12",
    "Atas/2026-04-02_Ata_riscos_ambientais.md": "2026-04-02",
    "Orcamentos/Orcamento_projeto_Lagoa_Norte.xlsx": "2026-03-20",
    "Relatorios/Nota_licenciamento_faixa_norte.md": "2026-04-08",
}


def _gravar(rel: str, dados: bytes) -> Path:
    destino = CORPUS / rel
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_bytes(dados)
    marca = MTIMES.get(rel)
    if marca:
        ts = datetime.strptime(marca, "%Y-%m-%d").replace(tzinfo=timezone.utc).timestamp()
        os.utime(destino, (ts, ts))
    return destino


def docx_politica(*, vigente: bool) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Política de Inteligência Artificial — Várzea Clara Energia", level=1)
    if vigente:
        d.add_paragraph(
            "Versão vigente, revisada em agosto de 2026. Esta é a norma em vigor. "
            "Substitui a série numerada PO-VCE-007_vN."
        )
        d.add_heading("Governança", level=2)
        d.add_paragraph(
            "O comitê deliberativo permanente aprova todo uso de modelo em processo "
            "produtivo. Não há mais comitê consultivo."
        )
        d.add_heading("Classificação de risco", level=2)
        d.add_paragraph(
            "Projetos de alto risco exigem revisão humana em todas as decisões que "
            "afetem pessoa identificável."
        )
    else:
        d.add_paragraph(
            "Versão 6, janeiro de 2025. Documento histórico. Não usar como norma vigente."
        )
        d.add_heading("Governança", level=2)
        d.add_paragraph(
            "O comitê consultivo emite parecer não vinculante. A diretoria decide depois."
        )
        d.add_heading("Classificação de risco", level=2)
        d.add_paragraph("Três níveis: baixo, médio, alto. Alto risco sobe à diretoria.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pptx_deck(*, versao: int) -> bytes:
    from pptx import Presentation

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = f"Governança de IA — Várzea Clara Energia (v{versao})"
    corpo = slide.placeholders[1]
    if versao == 2:
        corpo.text = (
            "Modelo vigente: um comitê deliberativo permanente.\n"
            "Esta é a versão 2 do deck, aprovada em setembro de 2025.\n"
            "A v1 ficou obsoleta e não deve ser apresentada a conselho."
        )
    else:
        corpo.text = (
            "Modelo antigo: três comitês consultivos rotativos.\n"
            "Versão 1 do deck. Superada pela v2."
        )
    slide.notes_slide.notes_text_frame.text = f"Deck interno VCE, versão {versao}."
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def pdf_texto(titulo: str, corpo: str) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    pagina = doc.new_page()
    pagina.insert_textbox(
        pymupdf.Rect(72, 64, 520, 120),
        titulo,
        fontsize=14,
        fontname="helv",
    )
    pagina.insert_textbox(
        pymupdf.Rect(72, 130, 520, 760),
        corpo,
        fontsize=11,
        fontname="helv",
    )
    dados = doc.tobytes()
    doc.close()
    return dados


def xlsx_orcamento() -> bytes:
    import openpyxl

    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Equipe"
    aba.append(["Papel", "Pessoas", "Custo mensal (BRL)"])
    aba.append(["Coordenação", 1, 8000])
    aba.append(["Analista de dados", 2, 12000])
    aba.append(["Total", 3, 20000])
    aba2 = livro.create_sheet("Alternativas")
    aba2.append(["Opção", "Descrição", "Custo mensal (BRL)"])
    aba2.append(["A", "Equipe interna mínima", 20000])
    aba2.append(["B", "Misto com dois consultores da Aurora Técnica", 34000])
    aba2.append(["C", "Sustentação só com a Boreal Serviços", 18000])
    buf = io.BytesIO()
    livro.save(buf)
    return buf.getvalue()


def gerar() -> list[Path]:
    if CORPUS.exists():
        shutil.rmtree(CORPUS)
    CORPUS.mkdir(parents=True)

    escritos = [
        _gravar(
            "Politica de IA/PO-VCE-007_Politica de Inteligencia Artificial_v6.docx",
            docx_politica(vigente=False),
        ),
        _gravar(
            "Politica de IA/PO-VCE-007_Politica de Inteligencia Artificial_revisada_GC.docx",
            docx_politica(vigente=True),
        ),
        _gravar(
            "Propostas/Proposta_Aurora_Tecnica_implantacao_IA.pdf",
            pdf_texto(
                "Proposta de implantação de IA — Aurora Técnica",
                "A Aurora Técnica propõe a implantação do programa de inteligência "
                "artificial da Várzea Clara Energia no Projeto Lagoa Norte. "
                "Escopo: classificação de documentos, fila de revisão humana e "
                "treinamento da equipe interna. Valor da proposta: R$ 480.000. "
                "Esta é a empresa que apresentou o plano de implantação, não de "
                "sustentação.",
            ),
        ),
        _gravar(
            "Propostas/Proposta_Boreal_Servicos_implantacao_IA.pdf",
            pdf_texto(
                "Proposta de sustentação de IA — Boreal Serviços",
                "A Boreal Serviços apresenta alternativa de sustentação operacional "
                "depois que o programa já estiver implantado. Não cobre a implantação "
                "inicial. Valor da proposta: R$ 210.000 por ano. Confundir esta "
                "proposta com a de implantação escolhe o fornecedor errado.",
            ),
        ),
        _gravar(
            "Contratos/CT-VCE-2024-0142_Servicos_consultoria.pdf",
            pdf_texto(
                "Contrato CT-VCE-2024-0142 — serviços de consultoria",
                "Contrato CT-VCE-2024-0142 firmado entre Várzea Clara Energia e o "
                "escritório Ribeira Advogados, em 3 de novembro de 2024, no valor "
                "de R$ 1.240.000, referente a serviços de consultoria jurídica do "
                "Projeto Lagoa Norte. Vigência de 24 meses. Código do processo "
                "interno: PRC-LGN-0142.",
            ),
        ),
        _gravar("Apresentacoes/Deck_governanca_IA_v1.pptx", pptx_deck(versao=1)),
        _gravar("Apresentacoes/Deck_governanca_IA_v2.pptx", pptx_deck(versao=2)),
        _gravar("Orcamentos/Orcamento_projeto_Lagoa_Norte.xlsx", xlsx_orcamento()),
        _gravar(
            "Atas/2026-03-12_Ata_mudanca_escopo.md",
            (
                "# Ata de 12/03/2026 — mudança de escopo\n\n"
                "Reunião do Projeto Lagoa Norte. Decisão: incluir a faixa norte "
                "no escopo da obra, além da margem já licenciada.\n\n"
                "A mudança entra em vigor na semana seguinte. Riscos ambientais "
                "ficam para a reunião de abril.\n"
            ).encode("utf-8"),
        ),
        _gravar(
            "Atas/2026-04-02_Ata_riscos_ambientais.md",
            (
                "# Ata de 02/04/2026 — riscos ambientais\n\n"
                "Depois da mudança de escopo de 12/03/2026, o grupo registrou "
                "dois riscos: atraso no licenciamento da faixa norte e "
                "interferência com nascente na margem esquerda.\n\n"
                "Encaminhamento: a nota técnica de licenciamento deve sair nesta "
                "semana.\n"
            ).encode("utf-8"),
        ),
        _gravar(
            "Relatorios/Nota_licenciamento_faixa_norte.md",
            (
                "# Nota técnica — autorização ambiental da margem esquerda\n\n"
                "A faixa norte do Projeto Lagoa Norte depende de autorização "
                "ambiental da margem esquerda do rio Clara. Sem esse despacho, "
                "a inclusão de escopo de março não pode ir a campo.\n\n"
                "O texto não usa a palavra licenciamento no título de propósito: "
                "é o caso semântico em que a pergunta e o documento não compartilham "
                "o termo.\n"
            ).encode("utf-8"),
        ),
    ]
    return escritos


def main() -> int:
    arquivos = gerar()
    print(f"{len(arquivos)} arquivos em {CORPUS}")
    for p in arquivos:
        print(f"  {p.relative_to(CORPUS)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
