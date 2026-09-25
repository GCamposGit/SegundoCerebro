"""Picture OCR stays off without a working GPU, and on when the test hooks it."""

from __future__ import annotations

import io

import pytest

from segundocerebro.ingest import ocr
from segundocerebro.ingest.parsers.slides import parse_pptx
from segundocerebro.ingest.raster_ocr import acrescentar_rasters
from segundocerebro.index.repesca import versao_efetiva

_GRAFICO_SEM_CACHE = """<?xml version="1.0" encoding="UTF-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">
  <c:chart><c:plotArea><c:barChart><c:ser>
    <c:tx><c:v>Licencas</c:v></c:tx>
    <c:val><c:numRef><c:f>Sheet1!$B$2:$B$3</c:f></c:numRef></c:val>
  </c:ser></c:barChart></c:plotArea></c:chart>
</c:chartSpace>
"""

_GRAFICO_COM_CACHE = """<?xml version="1.0" encoding="UTF-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">
  <c:chart><c:plotArea><c:barChart><c:ser>
    <c:tx><c:v>Licencas</c:v></c:tx>
    <c:cat><c:strCache><c:pt idx="0"><c:v>Jan</c:v></c:pt></c:strCache></c:cat>
    <c:val><c:numCache><c:pt idx="0"><c:v>10.5</c:v></c:pt></c:numCache></c:val>
  </c:ser></c:barChart></c:plotArea></c:chart>
</c:chartSpace>
"""


def bytes_pptx() -> bytes:
    from pptx import Presentation

    apresentacao = Presentation()
    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[1])
    slide.shapes.title.text = "Casos de uso do Copilot"
    slide.placeholders[1].text = "Resumir reunião\nRedigir e-mail"
    buf = io.BytesIO()
    apresentacao.save(buf)
    return buf.getvalue()


def _png(largura: int, altura: int) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (largura, altura), (240, 240, 240)).save(buf, format="PNG")
    return buf.getvalue()


def _zip_com(nome: str, bruto: bytes) -> bytes:
    import zipfile

    base = bytes_pptx()
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base), "r") as origem, zipfile.ZipFile(buf, "w") as destino:
        for item in origem.infolist():
            destino.writestr(item, origem.read(item.filename))
        destino.writestr(nome, bruto)
    return buf.getvalue()


@pytest.fixture
def motor(monkeypatch):
    chamadas: list[int] = []

    def falso(_imagem):
        chamadas.append(1)
        return "SCAN-VCE-001"

    monkeypatch.setattr(ocr, "motor_imagem", falso)
    monkeypatch.setattr(ocr, "_forcar_gpu", True)
    yield chamadas
    ocr.motor_imagem = None
    ocr._forcar_gpu = None
    ocr._gpu_cache = None
    ocr._rapid = None


def test_raster_entra_quando_o_gancho_esta_ligado(motor: list[int]) -> None:
    doc = parse_pptx(_zip_com("ppt/media/figura.png", _png(320, 180)), "deck.pptx")
    imagem = [b for b in doc.blocks if b.locator == "imagem"]
    assert len(imagem) == 1
    assert imagem[0].text == "SCAN-VCE-001"
    assert doc.meta.get("ocr_raster") == "gpu"
    assert motor


def test_raster_nao_repete_texto_que_o_slide_ja_tem(monkeypatch) -> None:
    monkeypatch.setattr(ocr, "motor_imagem", lambda _img: "Resumir reunião")
    monkeypatch.setattr(ocr, "_forcar_gpu", True)
    doc = parse_pptx(_zip_com("ppt/media/figura.png", _png(320, 180)), "deck.pptx")
    assert [b for b in doc.blocks if b.locator == "imagem"] == []
    ocr.motor_imagem = None
    ocr._forcar_gpu = None


def test_icone_nao_chama_o_motor(motor: list[int]) -> None:
    doc = parse_pptx(_zip_com("ppt/media/icone.png", _png(16, 16)), "deck.pptx")
    assert doc.meta.get("ocr_raster") != "gpu"
    assert motor == []


def test_filho_nao_dispara_ocr_de_imagem(monkeypatch) -> None:
    """O subprocesso não carrega CUDA. Sem isso a apresentação grande morre e sai do índice."""
    monkeypatch.setenv("SEGUNDOCEREBRO_SEM_OCR_RASTER", "1")
    monkeypatch.setattr(ocr, "motor_imagem", lambda _img: "SCAN-VCE-001")
    monkeypatch.setattr(ocr, "_forcar_gpu", True)
    meta = acrescentar_rasters(_zip_com("ppt/media/figura.png", _png(320, 180)), [])
    assert meta == {}
    ocr.motor_imagem = None
    ocr._forcar_gpu = None


def test_pai_acrescenta_o_raster_depois_do_filho(tmp_path, monkeypatch) -> None:
    from segundocerebro.ingest.document import ParseResult, ParseStatus
    from segundocerebro.index.isolamento import _raster_no_pai

    monkeypatch.setenv("SEGUNDOCEREBRO_SEM_OCR_RASTER", "1")
    doc = parse_pptx(_zip_com("ppt/media/figura.png", _png(320, 180)), "deck.pptx")
    assert all(b.locator != "imagem" for b in doc.blocks)
    caminho = tmp_path / "deck.pptx"
    caminho.write_bytes(_zip_com("ppt/media/figura.png", _png(320, 180)))
    monkeypatch.delenv("SEGUNDOCEREBRO_SEM_OCR_RASTER")
    monkeypatch.setattr(ocr, "motor_imagem", lambda _img: "SCAN-VCE-001")
    monkeypatch.setattr(ocr, "_forcar_gpu", True)
    saida = _raster_no_pai(
        str(caminho),
        ParseResult(path=str(caminho), status=ParseStatus.OK, doc=doc),
    )
    assert saida.doc is not None
    assert any(b.locator == "imagem" and b.text == "SCAN-VCE-001" for b in saida.doc.blocks)
    ocr.motor_imagem = None
    ocr._forcar_gpu = None


def test_sem_gpu_nao_chama_o_motor(monkeypatch) -> None:
    chamadas: list[int] = []
    monkeypatch.setattr(ocr, "motor_imagem", None)
    monkeypatch.setattr(ocr, "_forcar_gpu", False)

    def proibido(_img):
        chamadas.append(1)
        return "SCAN-VCE-001"

    monkeypatch.setattr(ocr, "_texto_rapidocr", proibido)
    meta = acrescentar_rasters(_zip_com("ppt/media/figura.png", _png(320, 180)), [])
    assert meta.get("ocr_raster") == "sem_gpu"
    assert chamadas == []
    ocr._forcar_gpu = None


def test_emf_fica_declarado(motor: list[int]) -> None:
    doc = parse_pptx(_zip_com("ppt/media/vetor.emf", b"nao-e-png"), "deck.pptx")
    assert "emf" in (doc.meta.get("imagem_nao_lida") or "")
    assert motor == []


def test_pastas_nvidia_le_a_raiz_indicada(tmp_path) -> None:
    from segundocerebro.index.cuda_runtime import pastas_nvidia

    binario = tmp_path / "nvidia" / "cudnn" / "bin"
    binario.mkdir(parents=True)
    assert pastas_nvidia([tmp_path]) == [str(binario)]


def test_planilha_embutida_supre_grafico_sem_cache() -> None:
    import zipfile

    import openpyxl

    livro = openpyxl.Workbook()
    livro.active.append(["VCE-PLANTA-9"])
    tabela = io.BytesIO()
    livro.save(tabela)
    base = _zip_com("ppt/charts/chart1.xml", _GRAFICO_SEM_CACHE.encode("utf-8"))
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base), "r") as origem, zipfile.ZipFile(buf, "w") as destino:
        for item in origem.infolist():
            destino.writestr(item, origem.read(item.filename))
        destino.writestr("ppt/embeddings/grafico.xlsx", tabela.getvalue())
    doc = parse_pptx(buf.getvalue(), "receita.pptx")
    texto = "\n".join(b.text for b in doc.blocks if b.locator == "grafico")
    assert "VCE-PLANTA-9" in texto


def test_planilha_embutida_nao_repete_valor_do_cache() -> None:
    import zipfile

    import openpyxl

    livro = openpyxl.Workbook()
    livro.active.append(["10.5"])
    tabela = io.BytesIO()
    livro.save(tabela)
    base = _zip_com("ppt/charts/chart1.xml", _GRAFICO_COM_CACHE.encode("utf-8"))
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base), "r") as origem, zipfile.ZipFile(buf, "w") as destino:
        for item in origem.infolist():
            destino.writestr(item, origem.read(item.filename))
        destino.writestr("ppt/embeddings/grafico.xlsx", tabela.getvalue())
    doc = parse_pptx(buf.getvalue(), "receita.pptx")
    grafico = "\n".join(b.text for b in doc.blocks if b.locator == "grafico")
    assert grafico.count("10.5") == 1


def test_versao_efetiva_so_marca_raster_com_gpu(monkeypatch) -> None:
    monkeypatch.setattr(ocr, "gpu_para_ocr", lambda: False)
    assert versao_efetiva(".pptx") == "3"
    assert versao_efetiva(".docx") == "2"
    monkeypatch.setattr(ocr, "gpu_para_ocr", lambda: True)
    assert versao_efetiva(".pptx") == "3+raster"
    assert versao_efetiva(".pdf") == "2"
