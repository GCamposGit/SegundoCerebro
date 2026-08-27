"""Tests for the parsers and for the single gatekeeper that opens files.

The most important guarantees here are not about extraction quality:

- no parser ever touches the filesystem (so the placeholder check can't be
  bypassed by whoever adds the next format);
- a locked file becomes a recorded status, not an exception;
- a scanned PDF is marked instead of being indexed as an empty document.
"""

from __future__ import annotations

import builtins
import io
import os
from pathlib import Path

import pytest

from segundocerebro.ingest import reader as reader_mod
from segundocerebro.ingest.document import BlockKind, ParseStatus
from segundocerebro.ingest.parsers import parser_for, supported_extensions
from segundocerebro.ingest.parsers.pdf import parse_pdf
from segundocerebro.ingest.parsers.sheets import parse_csv, parse_xls, parse_xlsx
from segundocerebro.ingest.parsers.slides import parse_ppt, parse_pptx
from segundocerebro.ingest.parsers.text import blocos_de_markdown, decode, parse_markdown, parse_rtf, parse_texto
from segundocerebro.ingest.parsers.word import _nivel, parse_docx
from segundocerebro.ingest.reader import CloudOnlyFile, FileLocked, parse_file, read_bytes

# --- construtores de arquivo sintético --------------------------------------


def bytes_docx(com_tabela: bool = False) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Política de IA", level=1)
    d.add_paragraph("Esta política define o uso aceitável.")
    d.add_heading("Classificação de risco", level=2)
    d.add_paragraph("Projetos são classificados em três níveis.")
    if com_tabela:
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Nível"
        t.cell(0, 1).text = "Ação"
        t.cell(1, 0).text = "Alto"
        t.cell(1, 1).text = "Human in the loop"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def bytes_pptx() -> bytes:
    from pptx import Presentation

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Casos de uso do Copilot"
    slide.placeholders[1].text = "Resumir reunião\nRedigir e-mail"
    slide.notes_slide.notes_text_frame.text = "Mencionar as licenças disponíveis."
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def bytes_xlsx(linhas: int = 70) -> bytes:
    import openpyxl

    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Licenças"
    aba.append(["Usuário", "Área", "Licença"])
    for i in range(linhas):
        aba.append([f"pessoa{i}", "Inovação", "Copilot M365"])
    buf = io.BytesIO()
    livro.save(buf)
    return buf.getvalue()


def bytes_pdf(texto: str | None = "Contrato 4600009999 com a Nimbus Tecnologia.", com_imagem: bool = False) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    pagina = doc.new_page()
    if com_imagem:
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 64))
        pix.set_rect(pix.irect, (200, 200, 200))
        pagina.insert_image(pymupdf.Rect(0, 0, 500, 700), pixmap=pix)
    if texto:
        pagina.insert_text((72, 72), texto, fontsize=11)
    dados = doc.tobytes()
    doc.close()
    return dados


# --- texto e markdown --------------------------------------------------------


def test_markdown_monta_a_trilha_de_titulos() -> None:
    md = b"# Plano\ntexto de abertura\n\n## Riscos\n### Ambientais\nlicenciamento pendente\n"
    doc = parse_markdown(md, "plano.md")

    trilhas = [b.heading_path for b in doc.blocks]
    assert ("Plano",) in trilhas
    assert ("Plano", "Riscos", "Ambientais") in trilhas
    ultimo = doc.blocks[-1]
    assert ultimo.text == "licenciamento pendente"
    assert ultimo.contextual_text.startswith("Plano > Riscos > Ambientais\n---\n")


def test_markdown_nao_confunde_cerca_de_codigo_com_titulo() -> None:
    md = b"# Doc\n```\n# isto e comentario, nao titulo\n```\nfim\n"
    doc = parse_markdown(md, "x.md")

    assert all(b.heading_path == ("Doc",) for b in doc.blocks)
    assert "# isto e comentario" in doc.blocks[0].text


def test_markdown_titulo_sublinhado() -> None:
    doc = parse_markdown(b"Introducao\n==========\nconteudo\n", "x.md")
    assert doc.blocks[0].heading_path == ("Introducao",)


def test_decode_aceita_cp1252() -> None:
    assert decode("Orçamento".encode("cp1252")) == "Orçamento"
    assert decode("Orçamento".encode("utf-8")) == "Orçamento"


def test_texto_puro_vira_um_bloco() -> None:
    doc = parse_texto("linha 1\nlinha 2".encode(), "notas.txt")
    assert len(doc.blocks) == 1
    assert doc.blocks[0].heading_path == ()


def test_pilha_compartilhada_atravessa_chamadas() -> None:
    pilha: list[tuple[int, str]] = []
    blocos_de_markdown("# Seção A\ninicio", "p. 1", pilha, fallback_titulo=False)
    seguintes = blocos_de_markdown("continuação na página seguinte", "p. 2", pilha, fallback_titulo=False)

    assert seguintes[0].heading_path == ("Seção A",)
    assert seguintes[0].locator == "p. 2"


# --- docx --------------------------------------------------------------------


@pytest.mark.parametrize(
    "estilo,esperado",
    [("Heading 1", 1), ("Heading 3", 3), ("Título 2", 2), ("Titulo 2", 2), ("Title", 1), ("Normal", None), (None, None)],
)
def test_nivel_de_estilo_em_dois_idiomas(estilo: str | None, esperado: int | None) -> None:
    assert _nivel(estilo) == esperado


def test_docx_hierarquia_de_titulos() -> None:
    doc = parse_docx(bytes_docx(), "politica.docx")

    trilhas = [b.heading_path for b in doc.blocks]
    assert ("Política de IA",) in trilhas
    assert ("Política de IA", "Classificação de risco") in trilhas


def test_docx_mantem_a_tabela_na_secao_certa() -> None:
    """Iterar paragraphs e tables em separado jogaria a tabela para o fim."""
    doc = parse_docx(bytes_docx(com_tabela=True), "politica.docx")

    tabelas = [b for b in doc.blocks if b.kind is BlockKind.TABLE]
    assert len(tabelas) == 1
    assert tabelas[0].heading_path == ("Política de IA", "Classificação de risco")
    assert "Human in the loop" in tabelas[0].text


# --- pptx --------------------------------------------------------------------


def test_pptx_slide_vira_bloco_com_titulo_e_notas() -> None:
    doc = parse_pptx(bytes_pptx(), "copilot.pptx")

    corpo = [b for b in doc.blocks if b.kind is BlockKind.TEXT]
    notas = [b for b in doc.blocks if b.kind is BlockKind.SLIDE_NOTES]

    assert corpo and corpo[0].heading_path[-1] == "Casos de uso do Copilot"
    assert corpo[0].locator == "slide 1"
    assert "Resumir reunião" in corpo[0].text
    assert notas and "licenças" in notas[0].text


# --- xlsx --------------------------------------------------------------------


def test_xlsx_repete_o_cabecalho_em_cada_janela() -> None:
    """Uma linha isolada não significa nada sem o cabeçalho."""
    doc = parse_xlsx(bytes_xlsx(linhas=70), "licencas.xlsx")

    assert len(doc.blocks) == 3  # 70 linhas em janelas de 30
    for bloco in doc.blocks:
        assert bloco.text.startswith("Usuário | Área | Licença")
        assert bloco.heading_path == ("Licenças",)
        assert bloco.kind is BlockKind.SHEET


def test_xlsx_locator_aponta_a_faixa() -> None:
    doc = parse_xlsx(bytes_xlsx(linhas=5), "licencas.xlsx")

    assert doc.blocks[0].locator == "Licenças!A2:C6"


def test_xlsx_aba_grande_usa_janela_grande_e_ganha_cartao() -> None:
    """Aba de despejo de dados: janela de 200 linhas + um cartão para achá-la.

    Com janela fixa de 30, as 48 planilhas da amostra do acervo geravam 4116
    blocos contra 2684 de 100 PDFs — planilha seria a maioria do índice.
    """
    doc = parse_xlsx(bytes_xlsx(linhas=250), "grande.xlsx")

    cartoes = [b for b in doc.blocks if "resumo da aba" in b.locator]
    janelas = [b for b in doc.blocks if "resumo da aba" not in b.locator]

    assert len(cartoes) == 1
    assert "250 linhas de dados" in cartoes[0].text
    assert "Usuário | Área | Licença" in cartoes[0].text
    assert len(janelas) == 2  # 250 linhas em janelas de 200
    assert janelas[0].locator == "Licenças!A2:C201"


def test_xlsx_aba_pequena_mantem_janela_fina_e_nao_ganha_cartao() -> None:
    """Aba de resumo é conteúdo: cada faixa merece precisão."""
    doc = parse_xlsx(bytes_xlsx(linhas=70), "pequena.xlsx")

    assert all("resumo da aba" not in b.locator for b in doc.blocks)
    assert len(doc.blocks) == 3  # janelas de 30


def bytes_xlsx_enorme(linhas: int = 6000) -> bytes:
    import openpyxl

    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Export"
    aba.append(["Fornecedor", "Contrato", "Valor"])
    for i in range(linhas):
        # três fornecedores distintos repetidos, contrato único por linha
        aba.append([f"FORN-{i % 3}", f"46000{i:05d}", i * 1.5])
    buf = io.BytesIO()
    livro.save(buf)
    return buf.getvalue()


def bytes_xlsx_larga(linhas: int = 1200, colunas: int = 22) -> bytes:
    """Aba abaixo do limiar de altura e muito acima do de área.

    A forma do arquivo real que produziu 3.277 chunks: 4.279 linhas × 22 colunas
    de dados de processo judicial. Aqui em escala menor, o suficiente para passar
    das 20 mil células.
    """
    import openpyxl

    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "resultado"
    aba.append([f"coluna_com_nome_longo_{c}" for c in range(colunas)])
    for i in range(linhas):
        aba.append([f"TRIB-{i % 5}" if c == 0 else f"valor-{i}-{c}" for c in range(colunas)])
    buf = io.BytesIO()
    livro.save(buf)
    return buf.getvalue()


def test_xlsx_aba_larga_vai_para_digesto_mesmo_sem_estourar_a_altura() -> None:
    """O defeito medido em 13/08/2026: altura sozinha não detecta despejo de dados.

    1.200 linhas está abaixo das 5.000 do limiar de altura, mas 1.200 × 22 passa
    das 20 mil células. Janelar renderia milhares de chunks quase idênticos, cada
    um com o cabeçalho longo repetido ocupando boa parte do vetor.

    A checagem por área roda depois da leitura de propósito: `max_row` e
    `max_column` em `read_only` vêm do elemento `<dimension>` do XML, que pode
    faltar — e faltou no arquivo real.
    """
    doc = parse_xlsx(bytes_xlsx_larga(1200, 22), "processos.xlsx")

    assert doc.meta.get("abas_em_digesto") == "resultado"
    assert all("valores distintos" in b.locator or "resumo da aba" in b.locator for b in doc.blocks), (
        "aba larga não pode cair no caminho de janelas"
    )
    tribunais = [b for b in doc.blocks if b.heading_path[-1:] == ("coluna_com_nome_longo_0",)]
    assert tribunais and tribunais[0].text.count("TRIB-") == 5, "digesto cobre a coluna categórica"


def test_xlsx_aba_de_conteudo_continua_janelada() -> None:
    """A checagem de área não pode engolir a aba de resumo, que é conteúdo.

    A aba mediana do acervo tem 230 linhas × 22 colunas ≈ 5.000 células, bem
    abaixo do limiar. Ali cada faixa de linhas merece precisão, e o digesto de
    valores distintos seria perda de informação.
    """
    doc = parse_xlsx(bytes_xlsx_larga(200, 20), "resumo.xlsx")

    assert "abas_em_digesto" not in doc.meta
    assert any("!A" in b.locator for b in doc.blocks), "faixa de linhas tem que aparecer no locator"


def test_xlsx_aba_enorme_cobre_coluna_categorica_ate_a_ultima_linha() -> None:
    """Coluna categórica fica COMPLETA, sem limite de linha.

    É o ganho real sobre o corte anterior em 5000 linhas: três fornecedores
    repetidos ao longo de 6000 linhas ficam inteiramente cobertos, e a linha
    6000 deixa de ser invisível para a busca.
    """
    doc = parse_xlsx(bytes_xlsx_enorme(6000), "export.xlsx")

    fornecedores = [b for b in doc.blocks if b.heading_path[-1:] == ("Fornecedor",)]
    assert len(fornecedores) == 1
    assert fornecedores[0].text.count("FORN-") == 3  # todos os distintos, dedup
    assert doc.meta["abas_em_digesto"] == "Export"
    assert all("valores distintos" in b.locator or "resumo da aba" in b.locator for b in doc.blocks)


def test_xlsx_aba_enorme_cobre_chave_unica_ate_o_teto_e_declara_o_limite() -> None:
    """Coluna de chave única não cabe inteira, e o limite tem que ser explícito.

    Com 9000 linhas de contrato único, o digesto cobre 8000 e registra a
    cobertura parcial no metadado. Fingir cobertura completa seria pior que
    declarar o teto.
    """
    doc = parse_xlsx(bytes_xlsx_enorme(9000), "export.xlsx")

    texto = "\n".join(b.text for b in doc.blocks)
    assert "4600000001" in texto  # início coberto
    assert "digesto_parcial" in doc.meta
    assert "Export!B" in doc.meta["digesto_parcial"]


def test_xlsx_digesto_ignora_coluna_de_medida() -> None:
    """Ninguém busca por '9000,0' — coluna decimal não é identificador."""
    doc = parse_xlsx(bytes_xlsx_enorme(6000), "export.xlsx")

    assert not [b for b in doc.blocks if b.heading_path[-1:] == ("Valor",)]


def test_xlsx_digesto_tem_cartao_unico() -> None:
    doc = parse_xlsx(bytes_xlsx_enorme(6000), "export.xlsx")

    cartao = [b for b in doc.blocks if "resumo da aba" in b.locator]
    assert len(cartao) == 1
    assert "6000 linhas" in cartao[0].text


def test_xlsx_digesto_para_de_crescer_com_o_tamanho_da_aba() -> None:
    """A propriedade que importa: custo limitado, não custo pequeno.

    6000 linhas de contrato único já custam ~40 blocos — é o preço de indexar
    uma coluna de chave única, e é comparável à janela por linha que havia
    antes. O ganho é que triplicar a aba não triplica o índice.
    """
    pequena = parse_xlsx(bytes_xlsx_enorme(6000), "a.xlsx")
    grande = parse_xlsx(bytes_xlsx_enorme(20000), "b.xlsx")

    assert len(grande.blocks) < 1.5 * len(pequena.blocks)
    assert len(grande.blocks) < 80


def test_csv_que_e_prosa_nao_perde_o_identificador() -> None:
    """O gerador planta `CT-FT-` em um `.csv` que é nota, não tabela."""
    bruto = "REGISTRO CT-FT-000\nContratada: Aurora.\nValor total: R$ 3.398.000.\n"
    doc = parse_csv(bruto.encode(), "Registro CT-FT-000.csv")
    assert "CT-FT-000" in "\n".join(b.text for b in doc.blocks)


def test_csv_repete_cabecalho_em_cada_janela() -> None:
    """A rota de texto deixava as janelas seguintes órfãs, sem esquema."""
    linhas = ["fornecedor,valor"] + [f"acme-{i},{i}" for i in range(80)]
    doc = parse_csv("\n".join(linhas).encode(), "compras.csv")

    janelas = [b for b in doc.blocks if "resumo" not in b.locator]
    assert len(janelas) >= 2
    assert all("fornecedor" in b.text.split("\n")[0].lower() for b in janelas)
    assert parser_for(".csv") is parse_csv


def test_csv_enorme_vira_digesto_nao_centenas_de_blobs() -> None:
    """Dump de 100k células: cartão + valores distintos, não 800 chunks de texto."""
    cab = "id,fornecedor,valor"
    # 7 000 × 3 > 20 000 células → digesto, o mesmo limiar da aba enorme.
    linhas = [cab] + [f"{i},acme-{i % 17},{i}.0" for i in range(7000)]
    doc = parse_csv("\n".join(linhas).encode(), "dump.csv")

    assert "abas_em_digesto" in doc.meta
    assert len(doc.blocks) < 80
    assert any("resumo da aba" in b.locator for b in doc.blocks)


def test_xlsx_vazio_nao_gera_bloco() -> None:
    import openpyxl

    buf = io.BytesIO()
    openpyxl.Workbook().save(buf)
    doc = parse_xlsx(buf.getvalue(), "vazia.xlsx")

    assert doc.blocks == ()
    assert "aviso" in doc.meta
    assert "sem_valor_em_cache" not in doc.meta


def bytes_xlsx_formula_sem_cache() -> bytes:
    """Labels plus `=SUM(1,2)` never opened in Excel — the C7.a trap."""
    import openpyxl

    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Orcamento"
    aba.append(["Rubrica", "Total apurado"])
    aba.append(["obras civis", "=SUM(1,2)"])
    buf = io.BytesIO()
    livro.save(buf)
    return buf.getvalue()


def bytes_xlsx_com_valor_calculado() -> bytes:
    """What LibreOffice is supposed to write: the number, not the formula."""
    import openpyxl

    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Orcamento"
    aba.append(["Rubrica", "Total apurado"])
    aba.append(["obras civis", 3])
    buf = io.BytesIO()
    livro.save(buf)
    return buf.getvalue()


def test_xlsx_formula_sem_cache_declara_e_nao_inventa_o_numero() -> None:
    """The class: the label is there, the Equity Value is not, and it is said."""
    doc = parse_xlsx(bytes_xlsx_formula_sem_cache(), "orcamento.xlsx")
    texto = "\n".join(b.text for b in doc.blocks)

    assert "obras civis" in texto
    assert "3" not in texto
    assert doc.meta.get("sem_valor_em_cache") == "1"


def test_xlsx_formula_sem_cache_recalcula_quando_soffice_devolve_valor(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The reader is the one that may spawn soffice; the parser still gets bytes."""
    monkeypatch.setattr(
        "segundocerebro.ingest.converters.libreoffice.recalcular_xlsx",
        lambda dados: bytes_xlsx_com_valor_calculado(),
    )
    alvo = tmp_path / "orcamento.xlsx"
    alvo.write_bytes(bytes_xlsx_formula_sem_cache())

    resultado = parse_file(str(alvo))
    texto = "\n".join(b.text for b in resultado.doc.blocks)

    assert resultado.status is ParseStatus.OK
    assert "obras civis" in texto
    assert "3" in texto
    assert resultado.doc.meta.get("recalculado") == "libreoffice"
    assert "sem_valor_em_cache" not in resultado.doc.meta


def test_xlsx_formula_sem_cache_sem_soffice_nao_some(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Missing binary is a visible warning, not EMPTY and not a crash."""
    monkeypatch.setattr(
        "segundocerebro.ingest.converters.libreoffice.recalcular_xlsx",
        lambda dados: None,
    )
    alvo = tmp_path / "orcamento.xlsx"
    alvo.write_bytes(bytes_xlsx_formula_sem_cache())

    resultado = parse_file(str(alvo))
    texto = "\n".join(b.text for b in resultado.doc.blocks)

    assert resultado.status is ParseStatus.OK
    assert "obras civis" in texto
    assert resultado.doc.meta.get("sem_valor_em_cache") == "1"


# --- pdf ---------------------------------------------------------------------


def test_pdf_extrai_texto_com_pagina_no_locator() -> None:
    doc = parse_pdf(bytes_pdf(), "contrato.pdf")

    assert doc.blocks
    assert "4600009999" in " ".join(b.text for b in doc.blocks)
    assert doc.blocks[0].locator == "p. 1"
    assert doc.meta["paginas"] == "1"


def bytes_pdf_estruturado() -> bytes:
    """PDF com título grande e corpo pequeno — o sinal que o motor 'fonte' usa."""
    import pymupdf

    doc = pymupdf.open()
    pagina = doc.new_page()
    pagina.insert_text((72, 80), "Relatório de Riscos", fontsize=22)
    pagina.insert_text((72, 130), "Ambientais", fontsize=16)
    y = 170
    for linha in (
        "O licenciamento da faixa norte depende da anuência do órgão estadual.",
        "O prazo estimado é de cento e oitenta dias a partir do protocolo.",
        "A equipe técnica acompanha o processo desde março de dois mil e vinte e seis.",
    ):
        pagina.insert_text((72, y), linha, fontsize=10)
        y += 20
    dados = doc.tobytes()
    doc.close()
    return dados


def test_pdf_motor_fonte_reconstroi_a_trilha_de_titulos() -> None:
    doc = parse_pdf(bytes_pdf_estruturado(), "riscos.pdf", motor="fonte")

    trilhas = [b.heading_path for b in doc.blocks if b.heading_path]
    assert ("Relatório de Riscos", "Ambientais") in trilhas
    corpo = [b for b in doc.blocks if "licenciamento" in b.text]
    assert corpo and corpo[0].heading_path == ("Relatório de Riscos", "Ambientais")
    assert doc.meta["motor"] == "fonte"


def test_pdf_motor_fonte_nao_confunde_paragrafo_longo_com_titulo() -> None:
    import pymupdf

    doc_pdf = pymupdf.open()
    pagina = doc_pdf.new_page()
    # texto longo em fonte grande: é destaque, não título
    pagina.insert_text((40, 80), "a" * 200, fontsize=20)
    pagina.insert_text((40, 200), "corpo do documento em fonte normal", fontsize=10)
    dados = doc_pdf.tobytes()
    doc_pdf.close()

    doc = parse_pdf(dados, "x.pdf", motor="fonte")

    assert all(len(b.heading_path) == 0 for b in doc.blocks)


def test_pdf_digitalizado_e_marcado_e_nao_indexado() -> None:
    doc = parse_pdf(bytes_pdf(texto=None, com_imagem=True), "escaneado.pdf")

    assert doc.blocks == ()
    assert doc.meta["suspeita"] == "digitalizado"


def test_pdf_pagina_de_imagem_com_pouco_texto_e_digitalizado() -> None:
    """Scan com carimbo de protocolo: tem imagem e um punhado de caracteres."""
    doc = parse_pdf(bytes_pdf(texto="Protocolo 4521/2025", com_imagem=True), "escaneado.pdf")

    assert doc.meta["suspeita"] == "digitalizado"


def test_pdf_curto_e_legitimo_nao_e_confundido_com_scan() -> None:
    """Recibo, fatura, passagem: pouco texto e nenhuma imagem — precisa indexar.

    O limiar inicial de 50 caracteres por página descartava esses documentos, e
    o acervo tem dezenas deles (14 faturas da OpenAI, recibos, passagens).
    """
    doc = parse_pdf(bytes_pdf(texto="Recibo 8871 — R$ 1.240,00"), "recibo.pdf")

    assert doc.blocks
    assert "suspeita" not in doc.meta


# --- registro ----------------------------------------------------------------


def test_registro_cobre_os_formatos_dominantes() -> None:
    for ext in (".pdf", ".xlsx", ".xls", ".docx", ".doc", ".pptx", ".ppt", ".md", ".txt", ".rtf"):
        assert parser_for(ext) is not None, ext
    assert parser_for(".dwg") is None
    assert ".pdf" in supported_extensions()


def test_extensao_e_case_insensitive() -> None:
    assert parser_for(".PDF") is parser_for(".pdf")


# --- leitor: as três armadilhas do corpus real -------------------------------


def test_nenhum_parser_abre_arquivo(monkeypatch: pytest.MonkeyPatch) -> None:
    """Parser recebe bytes. Se algum abrir caminho, o check de nuvem some."""

    # os construtores usam os templates das próprias bibliotecas, que estão em
    # disco — montar tudo antes de proibir a abertura
    docx_bytes, pptx_bytes, xlsx_bytes, pdf_bytes = (
        bytes_docx(),
        bytes_pptx(),
        bytes_xlsx(linhas=3),
        bytes_pdf(),
    )

    def proibido(*args, **kwargs):  # noqa: ANN002, ANN003
        raise AssertionError(f"parser tentou abrir arquivo: {args!r}")

    monkeypatch.setattr(builtins, "open", proibido)

    assert parse_docx(docx_bytes, "a.docx").blocks
    assert parse_pptx(pptx_bytes, "a.pptx").blocks
    assert parse_xlsx(xlsx_bytes, "a.xlsx").blocks
    assert parse_pdf(pdf_bytes, "a.pdf").blocks
    assert parse_markdown(b"# t\ncorpo", "a.md").blocks
    assert parse_ppt(pptx_bytes, "a.ppt").blocks
    assert parse_xls(
        b"<html><table><tr><td>CT-VCE-2024-0142</td></tr></table></html>", "a.xls"
    ).blocks


def test_placeholder_de_nuvem_e_recusado(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    alvo = tmp_path / "documento.docx"
    alvo.write_bytes(bytes_docx())
    monkeypatch.setattr(reader_mod, "is_cloud_only", lambda attrs: True)

    with pytest.raises(CloudOnlyFile):
        read_bytes(str(alvo))

    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.CLOUD_ONLY
    assert resultado.doc is None


def test_placeholder_pode_ser_hidratado_quando_explicito(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    alvo = tmp_path / "documento.md"
    alvo.write_bytes(b"# t\ncorpo")
    monkeypatch.setattr(reader_mod, "is_cloud_only", lambda attrs: True)

    assert read_bytes(str(alvo), allow_hydration=True) == b"# t\ncorpo"


def test_arquivo_travado_vira_status_nao_excecao(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Word segurando o arquivo é rotina neste corpus, não caso excepcional."""
    alvo = tmp_path / "plano.docx"
    alvo.write_bytes(bytes_docx())

    real = builtins.open

    def travado(caminho, *args, **kwargs):  # noqa: ANN001, ANN002, ANN003
        if str(caminho).endswith("plano.docx"):
            raise PermissionError(13, "usado por outro processo")
        return real(caminho, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", travado)

    with pytest.raises(FileLocked):
        read_bytes(str(alvo), retries=0)

    resultado = parse_file(str(alvo), retries=0)
    assert resultado.status is ParseStatus.LOCKED
    assert "outro processo" in resultado.detail


def test_travamento_transitorio_e_superado_na_segunda_tentativa(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    alvo = tmp_path / "nota.md"
    alvo.write_bytes(b"# t\ncorpo")

    real = builtins.open
    chamadas = {"n": 0}

    def as_vezes(caminho, *args, **kwargs):  # noqa: ANN001, ANN002, ANN003
        if str(caminho).endswith("nota.md"):
            chamadas["n"] += 1
            if chamadas["n"] == 1:
                raise PermissionError(13, "usado por outro processo")
        return real(caminho, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", as_vezes)

    assert read_bytes(str(alvo), retries=1, espera=0) == b"# t\ncorpo"


def test_caminho_acima_de_260_e_lido(tmp_path: Path) -> None:
    import shutil

    from segundocerebro.census import caminho_estendido

    fundo = tmp_path
    while len(str(fundo)) < 300:
        fundo = fundo / ("p" * 40)
    os.makedirs(caminho_estendido(str(fundo)), exist_ok=True)
    alvo = fundo / "documento.md"
    with open(caminho_estendido(str(alvo)), "wb") as fh:
        fh.write(b"# Longo\nconteudo")

    try:
        resultado = parse_file(str(alvo))
        assert resultado.ok
        assert resultado.doc is not None
        assert resultado.doc.blocks[0].heading_path == ("Longo",)
    finally:
        shutil.rmtree(caminho_estendido(str(tmp_path)), ignore_errors=True)


def test_formato_sem_parser(tmp_path: Path) -> None:
    alvo = tmp_path / "planta.dwg"
    alvo.write_bytes(b"binario")

    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.UNSUPPORTED
    assert resultado.detail == ".dwg"


def test_arquivo_corrompido_vira_erro_e_nao_derruba_a_ingestao(tmp_path: Path) -> None:
    alvo = tmp_path / "quebrado.docx"
    alvo.write_bytes(b"isto nao e um docx")

    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.ERROR
    assert resultado.doc is None


def test_documento_sem_texto_vira_vazio(tmp_path: Path) -> None:
    alvo = tmp_path / "escaneado.pdf"
    alvo.write_bytes(bytes_pdf(texto=None))

    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.EMPTY
    assert resultado.doc is not None


def test_rtf_extrai_texto_sem_abrir_arquivo() -> None:
    rtf = br"{\rtf1\ansi O contrato CT-VCE-2024-0142.\par }"
    doc = parse_rtf(rtf, "nota.rtf")
    assert "CT-VCE-2024-0142" in doc.blocks[0].text


def test_ppt_atomos_de_texto_sem_ole() -> None:
    import struct

    from segundocerebro.ingest.parsers.ole_texto import _ppt_registros

    payload = "Briefing VCE".encode("utf-16-le")
    atomo = struct.pack("<HHI", 0, 0x0FA0, len(payload)) + payload
    assert "Briefing VCE" in "\n".join(_ppt_registros(atomo))


# --- defeitos achados no acervo real, 15/08/2026 ------------------------------
#
# Os três vieram do checkpoint do run completo, não de imaginação sobre o que
# poderia dar errado. Cada teste guarda um caso que custou documento indexado.


def _com_validacao_de_coluna(dados: bytes) -> bytes:
    """Injeta `<dataValidations>` com `sqref="I:I"` — a faixa que o openpyxl recusa.

    Sem prefixo de namespace: o `openpyxl` gera a aba com namespace padrão, e um
    `x:` solto aqui daria erro de XML mal formado em vez do `TypeError` que este
    teste existe para reproduzir. Os arquivos reais usam prefixo porque o
    documento inteiro o declara — o que é indiferente para o openpyxl e está
    coberto por `test_filtro_de_validacoes_cobre_prefixo_de_namespace`.
    """
    import zipfile

    bloco = (
        b'<dataValidations count="1">'
        b'<dataValidation type="list" allowBlank="1" sqref="I:I">'
        b"<formula1>lista</formula1></dataValidation></dataValidations>"
    )
    saida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(dados)) as zin, zipfile.ZipFile(saida, "w") as zout:
        for item in zin.infolist():
            conteudo = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                conteudo = conteudo.replace(b"</worksheet>", bloco + b"</worksheet>")
            zout.writestr(item.filename, conteudo)
    return saida.getvalue()


def test_xlsx_com_validacao_de_coluna_inteira_e_recuperado() -> None:
    """`sqref="I:I"` é OOXML válido e o openpyxl recusa.

    Três arquivos do acervo caíam em `erro` por isso, um deles com 91 mil
    caracteres de conteúdo. A nova tentativa remove as validações — que não são
    indexadas — e recupera o documento.
    """
    from segundocerebro.ingest.parsers.sheets import _parse_xlsx

    dados = _com_validacao_de_coluna(bytes_xlsx(linhas=5))

    with pytest.raises(TypeError, match="MultiCellRange"):
        _parse_xlsx(dados, "orcamento.xlsx")

    doc = parse_xlsx(dados, "orcamento.xlsx")
    assert doc.blocks and "pessoa0" in doc.blocks[0].text


def test_xlsx_saudavel_nao_paga_a_nova_tentativa() -> None:
    """Sem validações não há o que reescrever — o caminho caro fica de fora."""
    from segundocerebro.ingest.parsers.sheets import _sem_validacoes

    assert _sem_validacoes(bytes_xlsx(linhas=3)) is None


def test_filtro_de_validacoes_cobre_prefixo_de_namespace() -> None:
    """Os arquivos do acervo escrevem `<x:dataValidations>`, não `<dataValidations>`.

    Procurar a forma sem prefixo no XML real não acha nada — foi o que atrasou o
    diagnóstico em 15/08/2026.
    """
    from segundocerebro.ingest.parsers.sheets import VALIDACOES

    com_prefixo = b'<x:sheetData/><x:dataValidations count="2"><x:dataValidation sqref="I:I"/></x:dataValidations><x:pageMargins/>'
    assert VALIDACOES.sub(b"", com_prefixo) == b"<x:sheetData/><x:pageMargins/>"

    sem_prefixo = b'<sheetData/><dataValidations count="1"><dataValidation sqref="A:A"/></dataValidations><pageMargins/>'
    assert VALIDACOES.sub(b"", sem_prefixo) == b"<sheetData/><pageMargins/>"


def test_pptx_com_forma_sem_frame_nao_derruba_o_deck() -> None:
    """`has_text_frame` verdadeiro e `text_frame` nulo — 4 decks perdidos por isso.

    O python-pptx devolve `None` quando o `<txBody>` falta, e o guarda booleano
    não cobre. Um slide assim levava o arquivo inteiro para `erro`.
    """
    from segundocerebro.ingest.parsers.slides import _texto_de_forma, _texto_do_frame

    class FormaQuebrada:
        has_text_frame = True
        text_frame = None

    class NotasQuebradas:
        notes_text_frame = None

    assert _texto_do_frame(FormaQuebrada()) == ""
    assert _texto_do_frame(NotasQuebradas(), "notes_text_frame") == ""
    assert _texto_de_forma(FormaQuebrada()) == []


def test_pptx_normal_continua_lendo() -> None:
    doc = parse_pptx(bytes_pptx(), "deck.pptx")
    assert any("Copilot" in b.text for b in doc.blocks)


def test_arquivo_apagado_entre_a_varredura_e_a_leitura_vira_sumiu(tmp_path: Path) -> None:
    """Corpus vivo: 10 dos 18 "erros" do run completo eram isto.

    Enquanto caía em `erro`, o contador misturava "reconciliar" com "consertar o
    parser" — que pedem coisas opostas.
    """
    resultado = parse_file(str(tmp_path / "revisao_apagada.xlsx"))
    assert resultado.status is ParseStatus.GONE
    assert resultado.status is not ParseStatus.ERROR
    assert resultado.doc is None


# --- F4-L: extensão que mente (HTML / criptografado / codepage / PPTX) --------
#
# Medido na passada de legado: `.xls` que é HTML ou SpreadsheetML, `.xls`
# criptografado, `xlrd` recusando codepage, `.ppt` que é PPTX. Sem o tratamento,
# os dois primeiros derrubam o parser; o PPTX-como-ppt vira `sem_parser` porque
# `ooxml` não tem parser sem ambiguidade no despachante. Vocabulário: VCE.


def _biff(tipo: int, payload: bytes) -> bytes:
    import struct

    return struct.pack("<HH", tipo, len(payload)) + payload


def _xls_ole_com_registros(*registros: bytes) -> bytes:
    import struct

    from cfb import escrever_cfb

    bof = _biff(0x0809, struct.pack("<HHHHHH", 0x0600, 0x0005, 0x0DBB, 0x07CC, 0, 0))
    eof = _biff(0x000A, b"")
    return escrever_cfb({"Workbook": bof + b"".join(registros) + eof})


def test_html_salvo_como_xls_extrai_celulas() -> None:
    html = (
        b"<html><table>"
        b"<tr><td>Contrato</td><td>Valor</td></tr>"
        b"<tr><td>CT-VCE-2024-0142</td><td>148500</td></tr>"
        b"</table></html>"
    )
    doc = parse_xls(html, "orcamento.xls")
    texto = "\n".join(b.text for b in doc.blocks)
    assert "CT-VCE-2024-0142" in texto
    assert doc.meta["conteudo_real"] == "html"


def test_html_salvo_como_xls_nao_vira_sem_parser(tmp_path: Path) -> None:
    html = (
        b"<html><table><tr><td>PO-VCE-007</td></tr></table></html>"
    )
    alvo = tmp_path / "politica.xls"
    alvo.write_bytes(html)
    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.OK
    assert resultado.status is not ParseStatus.UNSUPPORTED
    assert resultado.doc is not None
    assert "PO-VCE-007" in resultado.doc.blocks[0].text


def test_spreadsheetml_salvo_como_xls_extrai_aba() -> None:
    xml = (
        '<?xml version="1.0"?>'
        '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"'
        ' xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
        '<Worksheet ss:Name="Orcamento"><Table>'
        "<Row><Cell><Data ss:Type=\"String\">Contrato</Data></Cell>"
        "<Cell><Data ss:Type=\"String\">Valor</Data></Cell></Row>"
        "<Row><Cell><Data ss:Type=\"String\">CT-VCE-2024-0142</Data></Cell>"
        "<Cell><Data ss:Type=\"Number\">148500</Data></Cell></Row>"
        "</Table></Worksheet></Workbook>"
    ).encode()
    doc = parse_xls(xml, "orcamento.xls")
    texto = "\n".join(b.text for b in doc.blocks)
    assert "CT-VCE-2024-0142" in texto
    assert doc.meta["conteudo_real"] == "xml"
    assert any("Orcamento" in (b.heading_path or ()) for b in doc.blocks) or any(
        "Orcamento" in b.text for b in doc.blocks
    )


def test_xls_criptografado_biff_vira_erro_com_detalhe(tmp_path: Path) -> None:
    import struct

    filepass = _biff(0x002F, struct.pack("<HHHH", 1, 0, 1, 0) + b"\x00" * 16)
    dados = _xls_ole_com_registros(filepass)
    alvo = tmp_path / "secreto.xls"
    alvo.write_bytes(dados)

    with pytest.raises(ValueError, match="criptografada"):
        parse_xls(dados, "secreto.xls")

    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.ERROR
    assert resultado.status is not ParseStatus.UNSUPPORTED
    assert "criptograf" in resultado.detail.lower()


def test_xls_ole_com_encrypted_package_vira_erro_com_detalhe(tmp_path: Path) -> None:
    from cfb import escrever_cfb

    dados = escrever_cfb({"EncryptedPackage": b"x" * 80, "EncryptionInfo": b"y" * 20})
    alvo = tmp_path / "agile.xls"
    alvo.write_bytes(dados)
    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.ERROR
    assert resultado.status is not ParseStatus.UNSUPPORTED
    assert "criptograf" in resultado.detail.lower()


def test_xls_assertionerror_sem_mensagem_vira_erro_com_detalhe(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Achado do notebook: xlrd levanta AssertionError vazio num .xls real."""
    import xlrd

    def recusa(*args, **kwargs):  # noqa: ANN002, ANN003
        raise AssertionError()

    monkeypatch.setattr(xlrd, "open_workbook", recusa)
    dados = _xls_ole_com_registros()
    with pytest.raises(ValueError, match="xlrd recusou"):
        parse_xls(dados, "quebrado.xls")

    alvo = tmp_path / "quebrado.xls"
    alvo.write_bytes(dados)
    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.ERROR
    assert resultado.status is not ParseStatus.UNSUPPORTED
    assert "xlrd recusou" in resultado.detail
    assert "sem mensagem" in resultado.detail


def test_xls_codepage_recusada_abre_com_cp1252() -> None:
    import struct

    codepage = _biff(0x0042, struct.pack("<H", 0xABCD))
    dados = _xls_ole_com_registros(codepage)
    doc = parse_xls(dados, "legado.xls")
    assert doc.meta["formato"] == "xls"


def test_pptx_salvo_como_ppt_nao_vira_sem_parser(tmp_path: Path) -> None:
    dados = bytes_pptx()
    doc = parse_ppt(dados, "deck.ppt")
    assert any("Copilot" in b.text for b in doc.blocks)
    assert doc.meta["conteudo_real"] == "pptx"

    alvo = tmp_path / "deck.ppt"
    alvo.write_bytes(dados)
    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.OK
    assert resultado.status is not ParseStatus.UNSUPPORTED
    assert resultado.natureza is not None
    assert resultado.natureza.extensao_mente


def test_ppt_que_nao_e_ole2_vira_erro_nao_sem_parser(tmp_path: Path) -> None:
    alvo = tmp_path / "lixo.ppt"
    alvo.write_bytes(b"isto nao e um ppt")
    resultado = parse_file(str(alvo))
    assert resultado.status is ParseStatus.ERROR
    assert resultado.status is not ParseStatus.UNSUPPORTED
    assert "OLE2" in resultado.detail
